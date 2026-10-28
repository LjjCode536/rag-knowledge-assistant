"""HTTP 主接缝:文档解析链路(T02 / M1)。

覆盖五种格式解析、失败状态与原因、落盘位置可配置;
替身上游:解析全程本地进行,不触任何模型服务。
"""

from __future__ import annotations

from io import BytesIO

from docx import Document as new_docx
from fastapi.testclient import TestClient

from tests.fake_pdf import build_pdf

DOCUMENTS_URL = "/api/v1/documents"


def _upload(client: TestClient, *, name: str, content: bytes, mime: str) -> dict:
    response = client.post(DOCUMENTS_URL, files={"file": (name, content, mime)})
    return response.json()


# ===== 五种格式:解析成功,元素含类型与页码 / 定位信息 =====


def test_upload_txt_parses_paragraphs_with_line_positions(client):
    body = _upload(
        client,
        name="notes.txt",
        content="第一段第一行\n第一段第二行\n\n第二段".encode(),
        mime="text/plain",
    )

    assert body["code"] == "OK"
    data = body["data"]
    assert data["status"] == "parsed"
    assert data["file_type"] == "txt"
    assert data["error_message"] is None
    assert data["element_count"] == 2

    elements = client.get(f"{DOCUMENTS_URL}/{data['id']}/elements").json()["data"]
    assert elements["document_id"] == data["id"]
    first, second = elements["elements"]
    assert first["element_type"] == "paragraph"
    assert first["content"] == "第一段第一行\n第一段第二行"
    assert first["line"] == 1
    assert second["content"] == "第二段"
    assert second["line"] == 4


def test_upload_md_extracts_title_list_table_code(client):
    markdown = (
        "# 部署指南\n"
        "\n"
        "按照下列步骤执行。\n"
        "\n"
        "- 第一步\n"
        "- 第二步\n"
        "\n"
        "| 参数 | 值 |\n"
        "| --- | --- |\n"
        "| 超时 | 30 |\n"
        "\n"
        "```python\n"
        "print('hi')\n"
        "```"
    )
    body = _upload(client, name="guide.md", content=markdown.encode(), mime="text/markdown")

    data = body["data"]
    assert data["status"] == "parsed"
    elements = client.get(f"{DOCUMENTS_URL}/{data['id']}/elements").json()["data"]["elements"]

    assert [element["element_type"] for element in elements] == [
        "title",
        "paragraph",
        "list",
        "list",
        "table",
        "code",
    ]
    assert elements[0]["content"] == "部署指南"
    assert elements[0]["heading_level"] == 1
    assert elements[1]["line"] == 3
    assert elements[2]["content"] == "第一步"
    assert elements[4]["content"] == "参数 | 值\n超时 | 30"
    assert elements[5]["content"] == "print('hi')"


def test_upload_html_extracts_structure_and_skips_script(client):
    html = (
        "<html><head><style>body{color:red}</style></head><body>"
        "<script>var tracking = 1;</script>"
        "<h1>退费政策</h1>"
        "<p>7 天内可退。</p>"
        "<ul><li>在线申请</li><li>邮寄材料</li></ul>"
        "<table><tr><th>渠道</th><th>时效</th></tr><tr><td>App</td><td>3 天</td></tr></table>"
        "</body></html>"
    )
    body = _upload(client, name="policy.htm", content=html.encode(), mime="text/html")

    data = body["data"]
    assert data["status"] == "parsed"
    assert data["file_type"] == "html"
    elements = client.get(f"{DOCUMENTS_URL}/{data['id']}/elements").json()["data"]["elements"]

    assert [element["element_type"] for element in elements] == [
        "title",
        "paragraph",
        "list",
        "list",
        "table",
    ]
    assert elements[0]["heading_level"] == 1
    assert elements[4]["content"] == "渠道 | 时效\nApp | 3 天"
    joined = "\n".join(element["content"] for element in elements)
    assert "tracking" not in joined and "color:red" not in joined  # 噪声标签不泄漏


def test_upload_docx_extracts_title_paragraph_list_table(client):
    buffer = BytesIO()
    document = new_docx()
    document.add_heading("入职手册", level=1)
    document.add_paragraph("欢迎加入公司。")
    document.add_paragraph("领取门禁卡。", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "事项"
    table.cell(0, 1).text = "地点"
    table.cell(1, 0).text = "报到"
    table.cell(1, 1).text = "前台"
    document.save(buffer)

    body = _upload(
        client,
        name="handbook.docx",
        content=buffer.getvalue(),
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    data = body["data"]
    assert data["status"] == "parsed"
    assert data["file_type"] == "docx"
    elements = client.get(f"{DOCUMENTS_URL}/{data['id']}/elements").json()["data"]["elements"]

    assert [element["element_type"] for element in elements] == [
        "title",
        "paragraph",
        "list",
        "table",
    ]
    assert elements[0]["content"] == "入职手册"
    assert elements[0]["heading_level"] == 1
    assert elements[3]["content"] == "事项 | 地点\n报到 | 前台"


def test_upload_pdf_parses_paragraphs_with_page_numbers(client):
    pdf = build_pdf(["Refund window is seven days.", "Contact support for help."])
    body = _upload(client, name="policy.pdf", content=pdf, mime="application/pdf")

    data = body["data"]
    assert data["status"] == "parsed"
    assert data["file_type"] == "pdf"
    assert data["page_count"] == 2
    assert data["element_count"] == 2

    elements = client.get(f"{DOCUMENTS_URL}/{data['id']}/elements").json()["data"]["elements"]
    assert [(element["page_number"], element["content"]) for element in elements] == [
        (1, "Refund window is seven days."),
        (2, "Contact support for help."),
    ]


# ===== 解析失败:failed 状态 + error_message 可查,不静默吞错 =====


def test_corrupt_pdf_lands_in_failed_state_with_reason(client):
    body = _upload(client, name="broken.pdf", content=b"this is not a pdf", mime="application/pdf")

    assert body["code"] == "OK"  # 上传本身成功
    data = body["data"]
    assert data["status"] == "failed"
    assert "PDF" in data["error_message"]
    assert data["failed_at"]

    fetched = client.get(f"{DOCUMENTS_URL}/{data['id']}").json()["data"]
    assert fetched["status"] == "failed"
    assert fetched["error_message"] == data["error_message"]
    assert client.get(f"{DOCUMENTS_URL}/{data['id']}/elements").json()["data"]["elements"] == []


def test_corrupt_docx_lands_in_failed_state_with_reason(client):
    body = _upload(client, name="broken.docx", content=b"not a zip", mime="application/msword")

    data = body["data"]
    assert data["status"] == "failed"
    assert "DOCX" in data["error_message"]


def test_unsupported_format_rejected_without_record(client):
    response = client.post(
        DOCUMENTS_URL, files={"file": ("archive.xyz", b"data", "application/octet-stream")}
    )

    assert response.status_code == 415
    body = response.json()
    assert body["code"] == "UNSUPPORTED_FORMAT"
    assert body["trace_id"] == response.headers["X-Request-ID"]


def test_unknown_document_returns_404_envelope(client):
    response = client.get(f"{DOCUMENTS_URL}/no-such-document")

    assert response.status_code == 404
    body = response.json()
    assert body["code"] == "NOT_FOUND"
    assert body["details"] == {"document_id": "no-such-document"}


# ===== 落盘位置可配置,运行时产物不入库 =====


def test_uploaded_file_stored_under_configured_dir(client, upload_dir):
    body = _upload(client, name="notes.txt", content="内容".encode(), mime="text/plain")

    document_id = body["data"]["id"]
    files = list(upload_dir.iterdir())
    assert len(files) == 1
    assert files[0].name == f"{document_id}.txt"
    assert files[0].read_bytes() == "内容".encode()
